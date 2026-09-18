import hashlib
import io
from dataclasses import dataclass
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class UnsupportedFileType(Exception):
    pass


@dataclass
class IngestedDocument:
    source: str
    text: str
    content_hash: str


def load_document(path: str) -> IngestedDocument:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileType(f"Unsupported file type: {suffix or '(none)'}")
    file_bytes = file_path.read_bytes()
    if suffix == ".pdf":
        text = _extract_pdf_text(file_bytes)
    else:
        text = _extract_docx_text(file_bytes)
    content_hash = hashlib.sha256(file_bytes).hexdigest()
    return IngestedDocument(source=file_path.name, text=text, content_hash=content_hash)


def _extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_docx_text(file_bytes: bytes) -> str:
    doc = DocxDocument(io.BytesIO(file_bytes))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)
